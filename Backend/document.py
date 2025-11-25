import os
from typing import List

from docx import Document as Docx
from pptx import Presentation
from PyPDF2 import PdfReader


def _cleanup_lines(lines: List[str]) -> str:
    return "\n".join(line.strip() for line in lines if line and line.strip())


def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".docx":
        doc = Docx(file_path)
        lines = [p.text for p in doc.paragraphs]
        return _cleanup_lines(lines)

    if ext in {".pptx", ".ppt"}:
        try:
            pres = Presentation(file_path)
        except Exception as exc:  # pragma: no cover - library specific
            raise ValueError(
                "PowerPoint file could not be opened. Please upload a .pptx file."
            ) from exc

        lines = [
            shape.text
            for slide in pres.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        ]
        return _cleanup_lines(lines)

    if ext == ".pdf":
        reader = PdfReader(file_path)
        lines = []
        for page in reader.pages:
            text = page.extract_text() or ""
            lines.append(text)
        return _cleanup_lines(lines)

    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    raise ValueError(f"Unsupported file type: {ext}")